#!/usr/bin/env python3
"""DIGI-TEXX branded PowerPoint helpers (python-pptx).

Usage:
    python pptx_brand.py --demo out.pptx [--font Arial]

Library use:
    from pptx_brand import new_deck, add_cover, add_content, add_section, add_closing
    prs = new_deck()
    add_cover(prs, "Deck Title", "Subtitle")
    add_content(prs, "01. Section", "Headline", ["Point one", "Point two"])
    prs.save("deck.pptx")

All sizes below are px on the 1920x1080 design grid (see references/slides.md).
On a 13.333in slide 1px = 0.5pt; font sizes are converted with fpx().

Requires: pip install python-pptx pillow
"""
import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from PIL import Image

ASSETS = Path(__file__).resolve().parent.parent / "assets"
LOGO_SQUARE = ASSETS / "logos" / "logo-square.png"
LOGO_WHITE = ASSETS / "logos" / "logo-horizontal-white.png"
MASCOT = ASSETS / "mascot-robot.png"
WAVE = ASSETS / "patterns" / "pattern-wave-mesh.png"

BLUE = RGBColor(0x25, 0x82, 0xD7)
NAVY = RGBColor(0x26, 0x38, 0x5D)
SKY = RGBColor(0x30, 0xAA, 0xE0)
INK = RGBColor(0x23, 0x1F, 0x20)
GRAY = RGBColor(0x6E, 0x6F, 0x72)
DARK = RGBColor(0x1E, 0x1D, 0x23)
LIGHT = RGBColor(0xF0, 0xF0, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAD_START, GRAD_END = "00ACEB", "004EBE"

# Design grid is 1920x1080 px; convert px to EMU for a 13.333in x 7.5in slide.
PX = 12192000 / 1920

BODY_FONT = "Roboto"
DISPLAY_FONT = "Chakra Petch"  # fallback for SVN-HemiHead; set to "SVN-HemiHead" if installed


def px(v):
    return Emu(int(v * PX))


def fpx(v):
    """Font size given in design px -> Pt (1920px wide grid == 960pt slide)."""
    return Pt(v / 2)


def _picture_fill(slide, path, x, y, w, h):
    """Place an image covering the box (x, y, w, h), center-cropped, never stretched."""
    iw, ih = Image.open(path).size
    pic = slide.shapes.add_picture(str(path), px(x), px(y), px(w), px(h))
    box, img = w / h, iw / ih
    if img > box:  # image wider than box: crop left/right
        c = (1 - box / img) / 2
        pic.crop_left = pic.crop_right = c
    elif img < box:  # image taller: crop top/bottom
        c = (1 - img / box) / 2
        pic.crop_top = pic.crop_bottom = c
    return pic


def _bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _rect(slide, x, y, w, h, color=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
    s.line.fill.background()
    s.shadow.inherit = False
    if color is not None:
        s.fill.solid()
        s.fill.fore_color.rgb = color
    return s


def _gradient(shape, start=GRAD_START, end=GRAD_END, angle=0):
    """Apply a linear gradient fill (angle 0 = left to right)."""
    spPr = shape._element.spPr
    for tag in ("a:solidFill", "a:gradFill", "a:noFill"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    from lxml import etree
    xml = (
        '<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" rotWithShape="1">'
        f'<a:gsLst><a:gs pos="0"><a:srgbClr val="{start}"/></a:gs>'
        f'<a:gs pos="100000"><a:srgbClr val="{end}"/></a:gs></a:gsLst>'
        f'<a:lin ang="{angle * 60000}" scaled="0"/></a:gradFill>'
    )
    grad = etree.fromstring(xml)
    prstGeom = spPr.find(qn("a:prstGeom"))
    prstGeom.addnext(grad)


def _text(slide, x, y, w, h, text, size, color=WHITE, bold=False, italic=False,
          font=None, align=PP_ALIGN.LEFT, upper=False):
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for k, line in enumerate(lines):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line.upper() if upper else line
        f = r.font
        f.size = fpx(size)
        f.bold = bold
        f.italic = italic
        f.name = font or BODY_FONT
        f.color.rgb = color
    return tb


def _chrome(slide, title, page, light_panel):
    """Header title, gradient bar, optional light panel, corner logo, page number."""
    _bg(slide, DARK)
    if light_panel:
        _rect(slide, 0, 185, 1920, 895, LIGHT)
    _text(slide, 96, 70, 1400, 90, title, 48, WHITE, bold=True, italic=True,
          font=DISPLAY_FONT, upper=True)
    _gradient(_rect(slide, 1530, 87, 390, 70))
    slide.shapes.add_picture(str(LOGO_SQUARE), px(96), px(952), height=px(72))
    _page_number(slide, page, INK if light_panel else WHITE)


def _page_number(slide, page, color):
    _text(slide, 1620, 960, 204, 70, f"{page:02d}", 36, color,
          bold=True, italic=True, font=DISPLAY_FONT, align=PP_ALIGN.RIGHT)


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    return prs


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_cover(prs, title, subtitle="", image=None):
    s = _blank(prs)
    _bg(s, DARK)
    s.shapes.add_picture(str(LOGO_SQUARE), px(96), px(96), height=px(120))
    if image:
        _picture_fill(s, image, 960, 0, 960, 1080)
    else:
        _gradient(_rect(s, 960, 0, 960, 1080), "0D1120", "2582D7", angle=45)
    _text(s, 96, 300, 840, 400, title, 150, SKY, bold=True, italic=True, font=DISPLAY_FONT)
    _rect(s, 96, 730, 24, 6, BLUE)
    _rect(s, 120, 730, 72, 6, WHITE)
    if subtitle:
        _text(s, 96, 770, 820, 140, subtitle, 32, WHITE)
    _text(s, 96, 960, 600, 60, "www.digi-texx.com", 30, WHITE, bold=True, italic=True)
    return s


def add_content(prs, section_title, headline, bullets, page=None, light=True):
    s = _blank(prs)
    page = page or len(prs.slides)
    _chrome(s, section_title, page, light)
    fg = INK if light else WHITE
    _text(s, 96, 250, 1728, 90, headline, 44, fg, bold=True)
    _text(s, 96, 360, 1728, 560, [f"▪  {b}" for b in bullets], 28, fg)
    if not light:
        s.shapes.add_picture(str(WAVE), 0, px(820), width=px(1920))
    return s


def add_section(prs, title, page=None, image=None):
    s = _blank(prs)
    _bg(s, DARK)
    page = page or len(prs.slides)
    if image:
        _picture_fill(s, image, 0, 0, 1920, 1080)
    _gradient(_rect(s, 1530, 87, 390, 70))
    _text(s, 96, 380, 1400, 320, title, 120, WHITE if image else SKY, bold=True, italic=True,
          font=DISPLAY_FONT, upper=True)
    s.shapes.add_picture(str(LOGO_SQUARE), px(96), px(952), height=px(72))
    _page_number(s, page, WHITE)
    return s


def add_closing(prs, note="Let's get in touch — www.digi-texx.com", image=None):
    s = _blank(prs)
    if image:
        _picture_fill(s, image, 0, 0, 1920, 1080)
    else:
        _gradient(_rect(s, 0, 0, 1920, 1080), "0B4FA8", "2582D7", angle=90)
    _text(s, 600, 100, 1224, 260, "Thank You", 200, WHITE, bold=True, italic=True,
          font=DISPLAY_FONT, align=PP_ALIGN.RIGHT)
    s.shapes.add_picture(str(MASCOT), px(1000), px(380), height=px(130))
    _text(s, 1150, 400, 674, 120, note, 28, WHITE, italic=True)
    return s


def demo(out):
    prs = new_deck()
    add_cover(prs, "Deck Title", "Advanced digitizing services for your business",
              image=ASSETS / "patterns" / "pattern-circuit-left.jpg")
    add_content(prs, "01. Overview", "One clear message per slide",
                ["Lead with the business outcome", "Prove it with a real number",
                 "Close with a creative, direct CTA"])
    add_content(prs, "02. Approach", "Dark content slide with wave pattern",
                ["Use for statements and imagery"], light=False)
    add_section(prs, "Section Divider", image=ASSETS / "patterns" / "pattern-network-globe.jpg")
    add_closing(prs)
    prs.save(out)
    print(f"saved {out}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--demo", metavar="OUT", help="write a sample branded deck")
    ap.add_argument("--font", default=None, help="body font override, e.g. Arial for Vietnamese/Office decks")
    a = ap.parse_args()
    if a.font:
        BODY_FONT = a.font
    if a.demo:
        demo(a.demo)
    else:
        ap.print_help()
